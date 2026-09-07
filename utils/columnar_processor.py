"""
Processor for Multi-Column Hierarchical Decomposition / Marimekko Chart
with support for Excel Named Ranges (dataset-decomp.xlsx).
"""

from __future__ import annotations
import openpyxl
import re
import polars as pl
import pandas as pd
import numpy as np
from typing import List, Dict, Any, Optional, Tuple, Union

DEFAULT_DIMENSION_LABELS = {
    "Total TH": ["TOTAL"],
    "Total": ["TOTAL"],
    "Channel": ["SH CHAIN", "SH Online", "OT", "LMT", "CVS"],
    "Price": ["Tier1", "Tier2", "Tier3", "Tier4", "Tier5"],
    "Region": ["GBKK", "CENTRAL", "NORTH", "NORTHEAST", "SOUTH"],
    "Packsize": ["S", "M", "L", "XL", "XXL"]
}

DEFAULT_DIMENSION_NAMES = ["Total TH", "Channel", "Price", "Region", "Packsize"]

PRODUCT_ALIASES = {
    "dw": ["dish", "dishwash", "dishwashing", "dw"],
    "dish": ["dw", "dishwash", "dishwashing", "dish"],
    "dishwash": ["dw", "dish", "dishwashing"],
    "dishwashing": ["dw", "dish", "dishwash"],
    "fab_enh": ["fab_en", "fabric", "fabric_enhancer", "fab", "fab_enh"],
    "fab_en": ["fab_enh", "fabric", "fabric_enhancer", "fab", "fab_en"],
    "fabric": ["fab_enh", "fab_en", "fabric_enhancer", "fab"],
    "fabric_enhancer": ["fab_enh", "fab_en", "fabric", "fab"],
    "liquid": ["liq", "liquid"],
    "liq": ["liquid", "liq"],
    "powder": ["pwd", "pow", "powder"],
    "pwd": ["powder", "pow", "pwd"],
    "pow": ["powder", "pwd", "pow"],
}

DIMENSION_KEYWORDS = {
    "channel", "price", "region", "packsize", "pack", "tier", "total",
    "category", "brand", "dimension", "segment", "size"
}


def is_dimension_header_row(row_vals: List[Any]) -> bool:
    """Returns True if the row values contain dimension header keywords."""
    clean_vals = [str(v or "").strip().lower() for v in row_vals]
    matching = [v for v in clean_vals if any(k in v for k in DIMENSION_KEYWORDS)]
    return len(matching) >= 2


def classify_name(n: str) -> Tuple[Optional[str], Optional[str]]:
    """Classifies defined names into (role, token) where role in {'contribution', 'growth', 'bps', 'label'}."""
    nl = n.lower().strip()
    for p in ["contribution_", "cont_", "share_"]:
        if nl.startswith(p):
            return "contribution", nl[len(p):]
    for s in ["_contribution", "_cont", "_share"]:
        if nl.endswith(s):
            return "contribution", nl[:-len(s)]
    if nl.startswith("bps_"):
        return "bps", nl[4:]
    if nl.endswith("_bps"):
        return "bps", nl[:-4]
    for p in ["growth_pct_", "growth_", "gw_"]:
        if nl.startswith(p):
            return "growth", nl[len(p):]
    for s in ["_growth_pct", "_growth", "_gw"]:
        if nl.endswith(s):
            return "growth", nl[:-len(s)]
    for p in ["label_", "labels_", "dim_label_", "dim_labels_", "dim_"]:
        if nl.startswith(p):
            return "label", nl[len(p):]
    for s in ["_label", "_labels", "_dim"]:
        if nl.endswith(s):
            return "label", nl[:-len(s)]
    return None, None


def parse_range_str(ref_str: str):
    """Parses Excel reference string into (sheet, min_c, min_r, max_c, max_r)."""
    if not ref_str or "!" not in ref_str:
        return None, None, None, None, None
    from openpyxl.utils.cell import range_boundaries
    sheet_part, coord_part = ref_str.split("!", 1)
    clean_sheet = sheet_part.strip("'")
    clean_coord = coord_part.replace("$", "")
    if ":" not in clean_coord:
        clean_coord = f"{clean_coord}:{clean_coord}"
    min_c, min_r, max_c, max_r = range_boundaries(clean_coord)
    return clean_sheet, min_c, min_r, max_c, max_r


def is_numeric_value(val: Any) -> bool:
    """Checks if a cell value represents a numeric or metric quantity."""
    if val is None:
        return False
    s = str(val).strip()
    if not s or s.lower() == "none":
        return False
    s_clean = re.sub(r"[^\d\-+.]", "", s)
    if not s_clean:
        return False
    try:
        float(s_clean)
        return True
    except ValueError:
        return False


def clean_num(val: Any) -> Optional[float]:
    """Extracts clean float value from metric string (handles 'bps', '%', commas)."""
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() == "none":
        return None
    s_clean = re.sub(r"[^\d\-+.]", "", s)
    if not s_clean:
        return None
    try:
        return float(s_clean)
    except ValueError:
        return None


def parse_named_range_matrix(wb: openpyxl.Workbook, range_name: str) -> Tuple[List[str], List[List[Any]]]:
    """Extracts header and rows from an Excel defined name (openpyxl)."""
    if range_name not in wb.defined_names:
        return [], []
    from openpyxl.utils.cell import range_boundaries
    d = wb.defined_names[range_name]
    destinations = list(d.destinations) if hasattr(d, "destinations") else []
    formula = getattr(d, "value", "") or getattr(d, "attr_text", "")
    if not destinations and formula and "!" in formula:
        s_name, min_c, min_r, max_c, max_r = parse_range_str(formula)
        if s_name:
            destinations = [(s_name, f"{openpyxl.utils.get_column_letter(min_c)}{min_r}:{openpyxl.utils.get_column_letter(max_c)}{max_r}")]

    for title, coord in destinations:
        if not coord or coord == ":":
            continue
        sheet = wb[title]
        clean_coord = coord.replace("$", "")
        if ":" not in clean_coord:
            clean_coord = f"{clean_coord}:{clean_coord}"
        min_col, min_row, max_col, max_row = range_boundaries(clean_coord)

        first_row_vals = [sheet.cell(min_row, c).value for c in range(min_col, max_col + 1)]
        has_numeric = any(is_numeric_value(v) for v in first_row_vals)

        if has_numeric:
            header = []
            if min_row > 1:
                scan_vals = [str(sheet.cell(min_row - 1, c).value or "").strip() for c in range(min_col, max_col + 1)]
                if is_dimension_header_row(scan_vals) or (any(scan_vals) and not any(is_numeric_value(v) for v in scan_vals)):
                    header = scan_vals
            if not any(header):
                header = DEFAULT_DIMENSION_NAMES[: max_col - min_col + 1]
            matrix = [[sheet.cell(r, c).value for c in range(min_col, max_col + 1)] for r in range(min_row, max_row + 1)]
            return header, matrix
        else:
            header = [str(v or "").strip() for v in first_row_vals]
            matrix = [[sheet.cell(r, c).value for c in range(min_col, max_col + 1)] for r in range(min_row + 1, max_row + 1)]
            return header, matrix
    return [], []


def extract_dimension_labels(
    wb: openpyxl.Workbook,
    prod: str,
    cont_name: str,
    cont_headers: List[str]
) -> Dict[str, List[str]]:
    """Dynamically extracts dimension labels from workbook using defined names or adjacent columns."""
    from openpyxl.utils.cell import range_boundaries
    dynamic_labels: Dict[str, List[str]] = {}
    aliases = PRODUCT_ALIASES.get(prod.lower(), [prod.lower()])

    c_sheet, c_min_c, c_min_r, c_max_c, c_max_r = None, None, None, None, None
    if cont_name in wb.defined_names:
        c_dest = list(wb.defined_names[cont_name].destinations) if hasattr(wb.defined_names[cont_name], "destinations") else []
        c_val = getattr(wb.defined_names[cont_name], "value", "") or getattr(wb.defined_names[cont_name], "attr_text", "")
        if not c_dest and c_val and "!" in c_val:
            c_sheet, c_min_c, c_min_r, c_max_c, c_max_r = parse_range_str(c_val)
        elif c_dest and c_dest[0][1] and c_dest[0][1] != ":":
            c_sheet = c_dest[0][0]
            clean_c = c_dest[0][1].replace("$", "")
            if ":" not in clean_c:
                clean_c = f"{clean_c}:{clean_c}"
            c_min_c, c_min_r, c_max_c, c_max_r = range_boundaries(clean_c)

    # Strategy 1: Check label defined names by alias or matching row band
    for cand_name in list(wb.defined_names.keys()):
        role, token = classify_name(cand_name)
        if role != "label":
            continue

        d = wb.defined_names[cand_name]
        destinations = list(d.destinations) if hasattr(d, "destinations") else []
        formula = getattr(d, "value", "") or getattr(d, "attr_text", "")
        if not destinations and formula and "!" in formula:
            l_s, l_min_c, l_min_r, l_max_c, l_max_r = parse_range_str(formula)
            if l_s:
                destinations = [(l_s, f"{openpyxl.utils.get_column_letter(l_min_c)}{l_min_r}:{openpyxl.utils.get_column_letter(l_max_c)}{l_max_r}")]

        for title, coord in destinations:
            if not coord or coord == ":":
                continue
            sheet = wb[title]
            clean_coord = coord.replace("$", "")
            if ":" not in clean_coord:
                clean_coord = f"{clean_coord}:{clean_coord}"
            min_col, min_row, max_col, max_row = range_boundaries(clean_coord)

            # Check if row band matches or token matches
            row_match = (c_sheet and title.lower() == c_sheet.lower() and c_min_r is not None and max(c_min_r, min_row) <= min(c_max_r, max_row))
            token_match = (token == prod.lower() or token in aliases)

            if row_match or token_match:
                first_row_vals = [str(sheet.cell(min_row, c).value or "").strip() for c in range(min_col, max_col + 1)]
                has_headers = is_dimension_header_row(first_row_vals)
                data_start = min_row + 1 if has_headers else min_row

                num_cols = min(len(cont_headers), max_col - min_col + 1)
                for idx in range(num_cols):
                    target_key = cont_headers[idx]
                    dynamic_labels[target_key] = [
                        str(sheet.cell(r, min_col + idx).value or "").strip()
                        for r in range(data_start, max_row + 1)
                    ]
                if dynamic_labels:
                    return dynamic_labels

    # Strategy 2: Label table directly adjacent to the left
    if c_sheet and c_sheet in wb.sheetnames and c_min_c and c_min_c > 1:
        sheet = wb[c_sheet]
        adj_cols = min(len(cont_headers), c_min_c - 1)
        for idx in range(adj_cols):
            target_key = cont_headers[idx]
            dynamic_labels[target_key] = [
                str(sheet.cell(r, idx + 1).value or "").strip()
                for r in range(c_min_r, c_max_r + 1)
            ]
        if dynamic_labels:
            return dynamic_labels

    return dynamic_labels


import zipfile
import os
import io
import contextlib
import xml.etree.ElementTree as ET
try:
    import fastexcel
    HAS_FASTEXCEL = True
except ImportError:
    HAS_FASTEXCEL = False


@contextlib.contextmanager
def silence_rust_stderr():
    """Suppresses low-level C/Rust stderr messages like Calamine type inference warnings."""
    try:
        devnull = os.open(os.devnull, os.O_WRONLY)
        old_stderr = os.dup(2)
        os.dup2(devnull, 2)
        os.close(devnull)
        yield
    except Exception:
        yield
    finally:
        try:
            os.dup2(old_stderr, 2)
            os.close(old_stderr)
        except Exception:
            pass


def fast_extract_excel_named_ranges(file_source: Union[str, bytes, io.BytesIO]) -> Optional[Dict[str, Dict[str, Any]]]:
    """
    Ultra-fast selective named range and table extractor using raw zipfile parsing and fastexcel.
    Supports standard naming conventions: contribution_*, bps_*, growth_*, and label_*.
    """
    if not HAS_FASTEXCEL:
        return None

    try:
        from openpyxl.utils.cell import range_boundaries

        if isinstance(file_source, bytes):
            zf = zipfile.ZipFile(io.BytesIO(file_source))
            f_input = io.BytesIO(file_source)
        elif hasattr(file_source, "read"):
            b = file_source.read()
            zf = zipfile.ZipFile(io.BytesIO(b))
            f_input = io.BytesIO(b)
        else:
            zf = zipfile.ZipFile(file_source)
            f_input = file_source

        with zf:
            wb_xml = zf.read("xl/workbook.xml")
            root = ET.fromstring(wb_xml)
            ns = {"ns": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

            wb_rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
            rel_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in wb_rels}
            target_to_sheet_name = {}
            for s in root.findall(".//ns:sheet", ns):
                rId = s.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
                name = s.attrib["name"]
                target = rel_map.get(rId, "")
                if not target.startswith("worksheets/"):
                    target = "worksheets/" + target.split("/")[-1]
                target_to_sheet_name[target] = name

            sheet_offsets = {}
            for ws_path, ws_name in target_to_sheet_name.items():
                z_path = "xl/" + ws_path if not ws_path.startswith("xl/") else ws_path
                if z_path in zf.namelist():
                    ws_xml = zf.read(z_path)
                    m_dim = re.search(r'<dimension\s+ref="([^"]+)"', ws_xml.decode("utf-8", errors="ignore"))
                    if m_dim:
                        top_left = m_dim.group(1).split(":")[0]
                        b_c, b_r, _, _ = range_boundaries(f"{top_left}:{top_left}")
                        sheet_offsets[ws_name] = (b_c or 1, b_r or 1)
                    else:
                        sheet_offsets[ws_name] = (1, 1)

            classified_items = []

            # 1. Discover Excel Tables (ListObjects)
            sheet_file_to_name = {target.split("/")[-1]: name for target, name in target_to_sheet_name.items()}
            table_to_sheet = {}
            for zname in zf.namelist():
                if zname.startswith("xl/worksheets/_rels/") and zname.endswith(".rels"):
                    sheet_file = zname.replace("xl/worksheets/_rels/", "").replace(".rels", "")
                    sh_name = sheet_file_to_name.get(sheet_file)
                    if sh_name:
                        try:
                            rels_root = ET.fromstring(zf.read(zname))
                            for rel in rels_root:
                                if "table" in rel.attrib.get("Type", "").lower():
                                    tbl_target = rel.attrib["Target"].replace("\\", "/").split("/")[-1]
                                    table_to_sheet[tbl_target] = sh_name
                        except Exception:
                            pass

            for tbl_file, sh_name in table_to_sheet.items():
                tbl_path = "xl/tables/" + tbl_file
                if tbl_path in zf.namelist():
                    try:
                        t_root = ET.fromstring(zf.read(tbl_path))
                        t_name = t_root.get("name") or t_root.get("displayName")
                        t_ref = t_root.get("ref")
                        role, token = classify_name(t_name)
                        if role and t_ref:
                            min_c, min_r, max_c, max_r = range_boundaries(t_ref)
                            classified_items.append({
                                "name": t_name,
                                "role": role,
                                "token": token,
                                "sheet": sh_name,
                                "min_c": min_c,
                                "min_r": min_r,
                                "max_c": max_c,
                                "max_r": max_r,
                                "formula": t_ref
                            })
                    except Exception:
                        pass

            # 2. Discover Defined Names
            for dn in root.findall(".//ns:definedName", ns):
                name = dn.attrib.get("name")
                text = dn.text
                if name and text:
                    d_text = text.strip()
                    role, token = classify_name(name)
                    if role:
                        s_name, min_c, min_r, max_c, max_r = parse_range_str(d_text)
                        if s_name:
                            classified_items.append({
                                "name": name,
                                "role": role,
                                "token": token,
                                "sheet": s_name,
                                "min_c": min_c,
                                "min_r": min_r,
                                "max_c": max_c,
                                "max_r": max_r,
                                "formula": d_text
                            })

        with silence_rust_stderr():
            excel = fastexcel.read_excel(f_input)
        loaded_sheets = {}

        def get_sheet_df(sheet_name: str) -> pl.DataFrame:
            clean = sheet_name.strip("'")
            if clean not in loaded_sheets:
                with silence_rust_stderr():
                    s = excel.load_sheet_by_name(clean, header_row=None, dtype_coercion="coerce")
                    loaded_sheets[clean] = s.to_polars()
            return loaded_sheets[clean]

        cont_items = [i for i in classified_items if i["role"] == "contribution"]
        if not cont_items:
            cont_items = [i for i in classified_items if i["role"] == "growth"]

        if not cont_items:
            return None

        product_groups = []
        for c_item in cont_items:
            c_token = c_item["token"]
            c_sheet = c_item["sheet"]
            c_min_r = c_item["min_r"]
            c_max_r = c_item["max_r"]
            aliases = PRODUCT_ALIASES.get(c_token, [c_token])

            b_item = None
            for cand in classified_items:
                if cand["role"] == "bps":
                    if cand["sheet"].lower() == c_sheet.lower() and cand["min_r"] == c_min_r and cand["max_r"] == c_max_r:
                        b_item = cand
                        break
                    elif cand["token"] == c_token or cand["token"] in aliases:
                        b_item = cand
                        break

            g_item = None
            for cand in classified_items:
                if cand["role"] == "growth":
                    if cand["sheet"].lower() == c_sheet.lower() and cand["min_r"] == c_min_r and cand["max_r"] == c_max_r:
                        g_item = cand
                        break
                    elif cand["token"] == c_token or cand["token"] in aliases:
                        g_item = cand
                        break

            l_item = None
            for cand in classified_items:
                if cand["role"] == "label":
                    if cand["sheet"].lower() == c_sheet.lower() and cand["min_r"] == c_min_r and cand["max_r"] == c_max_r:
                        l_item = cand
                        break
            if not l_item:
                for cand in classified_items:
                    if cand["role"] == "label":
                        if cand["sheet"].lower() == c_sheet.lower() and max(cand["min_r"], c_min_r) <= min(cand["max_r"], c_max_r):
                            l_item = cand
                            break
                        elif cand["token"] == c_token or cand["token"] in aliases:
                            l_item = cand
                            break

            product_groups.append({
                "token": c_token,
                "cont": c_item,
                "bps": b_item,
                "growth": g_item,
                "label": l_item
            })

        sheet_master_headers = {}
        for ws_name in set(ci["sheet"] for ci in classified_items):
            try:
                s_df = get_sheet_df(ws_name)
                s_off_c, s_off_r = sheet_offsets.get(ws_name, (1, 1))
                for r_idx in range(min(10, s_df.shape[0])):
                    r_vals = [str(s_df.row(r_idx)[c] or "").strip() for c in range(s_df.shape[1])]
                    if is_dimension_header_row(r_vals):
                        sheet_master_headers[ws_name.lower()] = {
                            "row_idx": r_idx,
                            "vals": r_vals,
                            "off_c": s_off_c
                        }
                        break
            except Exception:
                pass

        results = {}

        for grp in product_groups:
            token = grp["token"]
            c_item = grp["cont"]
            b_item = grp["bps"]
            g_item = grp["growth"]
            l_item = grp["label"]

            sheet_df = get_sheet_df(c_item["sheet"])
            off_c, off_r = sheet_offsets.get(c_item["sheet"], (1, 1))

            c_min_c, c_min_r = c_item["min_c"], c_item["min_r"]
            c_max_c, c_max_r = c_item["max_c"], c_item["max_r"]

            df_row_min = c_min_r - off_r
            num_cols = c_max_c - c_min_c + 1
            cols_df_idx = [c_min_c - off_c + i for i in range(num_cols)]

            row_at_min_vals = [sheet_df.row(df_row_min)[ci] for ci in cols_df_idx]
            has_numeric_at_min = any(is_numeric_value(v) for v in row_at_min_vals)

            if has_numeric_at_min:
                data_start_r = c_min_r
                data_end_r = c_max_r
                headers = []
                if c_min_r > off_r:
                    scan_vals = [str(sheet_df.row(c_min_r - 1 - off_r)[ci] or "").strip() for ci in cols_df_idx]
                    if is_dimension_header_row(scan_vals) or (any(scan_vals) and not any(is_numeric_value(v) for v in scan_vals)):
                        headers = scan_vals

                if not any(headers):
                    m_hdr = sheet_master_headers.get(c_item["sheet"].lower())
                    if m_hdr:
                        m_vals = m_hdr["vals"]
                        m_off_c = m_hdr["off_c"]
                        m_slice = []
                        for ci in range(c_min_c, c_max_c + 1):
                            idx = ci - m_off_c
                            m_slice.append(m_vals[idx] if 0 <= idx < len(m_vals) else "")
                        if any(m_slice):
                            headers = m_slice

                if not any(headers):
                    headers = DEFAULT_DIMENSION_NAMES[:num_cols]
            else:
                headers = [str(v or "").strip() for v in row_at_min_vals]
                data_start_r = c_min_r + 1
                data_end_r = c_max_r

            target_N = data_end_r - data_start_r + 1

            cont_m = []
            for r in range(data_start_r, data_end_r + 1):
                df_r = r - off_r
                row_vals = [sheet_df.row(df_r)[ci] for ci in cols_df_idx]
                cont_m.append(row_vals)

            bps_m = []
            if b_item:
                b_df = get_sheet_df(b_item["sheet"])
                b_off_c, b_off_r = sheet_offsets.get(b_item["sheet"], (1, 1))
                b_cols_idx = [b_item["min_c"] - b_off_c + i for i in range(num_cols)]
                b_tot_r = b_item["max_r"] - b_item["min_r"] + 1
                b_start = b_item["min_r"] if b_tot_r == target_N else (b_item["min_r"] + 1)
                b_end = b_item["max_r"]
                for r in range(b_start, b_end + 1):
                    df_r = r - b_off_r
                    bps_m.append([b_df.row(df_r)[ci] for ci in b_cols_idx])

            gw_m = []
            if g_item:
                g_df = get_sheet_df(g_item["sheet"])
                g_off_c, g_off_r = sheet_offsets.get(g_item["sheet"], (1, 1))
                g_cols_idx = [g_item["min_c"] - g_off_c + i for i in range(num_cols)]
                g_tot_r = g_item["max_r"] - g_item["min_r"] + 1
                g_start = g_item["min_r"] if g_tot_r == target_N else (g_item["min_r"] + 1)
                g_end = g_item["max_r"]
                for r in range(g_start, g_end + 1):
                    df_r = r - g_off_r
                    gw_m.append([g_df.row(df_r)[ci] for ci in g_cols_idx])

            labels_by_col_idx = {}
            category_display_name = token.upper()

            if l_item:
                l_df = get_sheet_df(l_item["sheet"])
                l_off_c, l_off_r = sheet_offsets.get(l_item["sheet"], (1, 1))
                l_num_cols = l_item["max_c"] - l_item["min_c"] + 1
                l_cols_idx = [l_item["min_c"] - l_off_c + i for i in range(l_num_cols)]
                l_tot_r = l_item["max_r"] - l_item["min_r"] + 1
                l_start = l_item["min_r"] if l_tot_r == target_N else (l_item["min_r"] + 1)
                l_end = l_item["max_r"]

                label_matrix = []
                for r in range(l_start, l_end + 1):
                    df_r = r - l_off_r
                    label_matrix.append([str(l_df.row(df_r)[ci] or "").strip() for ci in l_cols_idx])

                if label_matrix and label_matrix[0] and label_matrix[0][0] and label_matrix[0][0].lower() != "none":
                    category_display_name = label_matrix[0][0]

                for ci in range(min(num_cols, l_num_cols)):
                    labels_by_col_idx[ci] = [label_matrix[ri][ci] for ri in range(len(label_matrix))]
            else:
                if c_min_c > 1:
                    adj_col_start = 1
                    adj_cols_count = min(num_cols, c_min_c - 1)
                    adj_cols_idx = [adj_col_start - off_c + i for i in range(adj_cols_count)]
                    for idx, ci in enumerate(adj_cols_idx):
                        col_vals = [str(sheet_df.row(r - off_r)[ci] or "").strip() for r in range(data_start_r, data_end_r + 1)]
                        labels_by_col_idx[idx] = col_vals
                    if 0 in labels_by_col_idx and labels_by_col_idx[0] and labels_by_col_idx[0][0]:
                        v0 = labels_by_col_idx[0][0]
                        if v0 and v0.lower() != "none":
                            category_display_name = v0.upper()

            columns_data = []
            flat_records = []

            for col_idx, col_name in enumerate(headers):
                if not col_name:
                    col_name = DEFAULT_DIMENSION_NAMES[col_idx] if col_idx < len(DEFAULT_DIMENSION_NAMES) else f"Dim_{col_idx+1}"

                dim_labels = labels_by_col_idx.get(col_idx, [])
                if not dim_labels:
                    dim_labels = DEFAULT_DIMENSION_LABELS.get(col_name, [])

                blocks = []
                for row_idx in range(len(cont_m)):
                    share_raw = cont_m[row_idx][col_idx] if row_idx < len(cont_m) and col_idx < len(cont_m[row_idx]) else None
                    share_num = clean_num(share_raw)
                    if share_num is None:
                        continue

                    share_pct = round(share_num * 100.0, 1) if 0 < share_num <= 1.0 else round(share_num, 1)

                    bps_num = 0
                    if bps_m and row_idx < len(bps_m) and col_idx < len(bps_m[row_idx]):
                        b_raw = bps_m[row_idx][col_idx]
                        b_clean = clean_num(b_raw)
                        bps_num = int(round(b_clean)) if b_clean is not None else 0

                    gw_pct = 0.0
                    if gw_m and row_idx < len(gw_m) and col_idx < len(gw_m[row_idx]):
                        g_raw = gw_m[row_idx][col_idx]
                        g_clean = clean_num(g_raw)
                        if g_clean is not None:
                            gw_pct = round(g_clean * 100.0, 1) if abs(g_clean) <= 5.0 and g_clean != 0 else round(g_clean, 1)

                    if col_idx == 0 or "total" in col_name.lower():
                        actual_label = category_display_name
                        generic_label = category_display_name
                    elif dim_labels and row_idx < len(dim_labels) and dim_labels[row_idx] and dim_labels[row_idx].lower() != "none":
                        actual_label = dim_labels[row_idx]
                        generic_label = f"ITEM {row_idx + 1}"
                    else:
                        actual_label = f"ITEM {row_idx + 1}"
                        generic_label = f"ITEM {row_idx + 1}"

                    is_pos = (bps_num >= 0 if b_item else gw_pct >= 0)
                    blocks.append({
                        "label": actual_label,
                        "actual_label": actual_label,
                        "generic_label": generic_label,
                        "share": share_pct,
                        "growth": gw_pct,
                        "bps": bps_num,
                        "is_positive": is_pos
                    })
                    flat_records.append({
                        "Product": category_display_name,
                        "Column": col_name,
                        "Label": actual_label,
                        "Share %": share_pct,
                        "Growth %": gw_pct,
                        "BPS Impact": bps_num
                    })

                columns_data.append({
                    "column_name": col_name,
                    "blocks": blocks,
                    "total_share": sum(b["share"] for b in blocks)
                })

            results[category_display_name] = {
                "columns_data": columns_data,
                "df": pl.DataFrame(flat_records)
            }

        return results
    except Exception:
        return None


def openpyxl_extract_excel_named_ranges(file_source: Union[str, bytes, io.BytesIO]) -> Optional[Dict[str, Dict[str, Any]]]:
    """Graceful openpyxl fallback for named range parsing."""
    try:
        from openpyxl.utils.cell import range_boundaries

        if isinstance(file_source, bytes):
            wb_input = io.BytesIO(file_source)
        else:
            wb_input = file_source

        wb = openpyxl.load_workbook(wb_input, data_only=True)
        all_names = list(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else [d.name for d in wb.defined_names.definedName]

        classified_items = []
        for name in all_names:
            d = wb.defined_names[name]
            destinations = list(d.destinations) if hasattr(d, "destinations") else []
            formula = getattr(d, "value", "") or getattr(d, "attr_text", "")
            if not destinations and formula and "!" in formula:
                s_name, min_c, min_r, max_c, max_r = parse_range_str(formula)
                if s_name:
                    destinations = [(s_name, f"{openpyxl.utils.get_column_letter(min_c)}{min_r}:{openpyxl.utils.get_column_letter(max_c)}{max_r}")]

            for sheet_title, coord in destinations:
                if not coord or coord == ":":
                    continue
                role, token = classify_name(name)
                if role:
                    coord_clean = coord.replace("$", "")
                    if ":" not in coord_clean:
                        coord_clean = f"{coord_clean}:{coord_clean}"
                    min_c, min_r, max_c, max_r = range_boundaries(coord_clean)
                    classified_items.append({
                        "name": name,
                        "role": role,
                        "token": token,
                        "sheet": sheet_title,
                        "min_c": min_c,
                        "min_r": min_r,
                        "max_c": max_c,
                        "max_r": max_r
                    })

        # Also discover Excel Tables (ListObjects) across all worksheets
        for ws in wb.worksheets:
            if hasattr(ws, "tables"):
                for t_name in ws.tables:
                    role, token = classify_name(t_name)
                    if role:
                        tbl = ws.tables[t_name]
                        t_ref = getattr(tbl, "ref", None)
                        if t_ref:
                            min_c, min_r, max_c, max_r = range_boundaries(t_ref)
                            classified_items.append({
                                "name": t_name,
                                "role": role,
                                "token": token,
                                "sheet": ws.title,
                                "min_c": min_c,
                                "min_r": min_r,
                                "max_c": max_c,
                                "max_r": max_r
                            })

        cont_items = [i for i in classified_items if i["role"] == "contribution"]
        if not cont_items:
            cont_items = [i for i in classified_items if i["role"] == "growth"]

        if not cont_items:
            return None

        product_groups = []
        for c_item in cont_items:
            c_token = c_item["token"]
            c_sheet = c_item["sheet"]
            c_min_r = c_item["min_r"]
            c_max_r = c_item["max_r"]
            aliases = PRODUCT_ALIASES.get(c_token, [c_token])

            b_item = None
            for cand in classified_items:
                if cand["role"] == "bps":
                    if cand["sheet"].lower() == c_sheet.lower() and cand["min_r"] == c_min_r and cand["max_r"] == c_max_r:
                        b_item = cand
                        break
                    elif cand["token"] == c_token or cand["token"] in aliases:
                        b_item = cand
                        break

            g_item = None
            for cand in classified_items:
                if cand["role"] == "growth":
                    if cand["sheet"].lower() == c_sheet.lower() and cand["min_r"] == c_min_r and cand["max_r"] == c_max_r:
                        g_item = cand
                        break
                    elif cand["token"] == c_token or cand["token"] in aliases:
                        g_item = cand
                        break

            l_item = None
            for cand in classified_items:
                if cand["role"] == "label":
                    if cand["sheet"].lower() == c_sheet.lower() and cand["min_r"] == c_min_r and cand["max_r"] == c_max_r:
                        l_item = cand
                        break
            if not l_item:
                for cand in classified_items:
                    if cand["role"] == "label":
                        if cand["sheet"].lower() == c_sheet.lower() and max(cand["min_r"], c_min_r) <= min(cand["max_r"], c_max_r):
                            l_item = cand
                            break
                        elif cand["token"] == c_token or cand["token"] in aliases:
                            l_item = cand
                            break

            product_groups.append({
                "token": c_token,
                "cont": c_item,
                "bps": b_item,
                "growth": g_item,
                "label": l_item
            })

        sheet_master_headers = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for r_idx in range(1, min(10, ws.max_row + 1)):
                r_vals = [str(ws.cell(r_idx, c).value or "").strip() for c in range(1, min(35, ws.max_column + 1))]
                if is_dimension_header_row(r_vals):
                    sheet_master_headers[sheet_name.lower()] = {
                        "row_idx": r_idx,
                        "vals": r_vals
                    }
                    break

        results = {}

        for grp in product_groups:
            token = grp["token"]
            c_item = grp["cont"]
            b_item = grp["bps"]
            g_item = grp["growth"]
            l_item = grp["label"]

            ws = wb[c_item["sheet"]]
            c_min_c, c_min_r = c_item["min_c"], c_item["min_r"]
            c_max_c, c_max_r = c_item["max_c"], c_item["max_r"]
            num_cols = c_max_c - c_min_c + 1

            row_at_min_vals = [ws.cell(c_min_r, ci).value for ci in range(c_min_c, c_max_c + 1)]
            has_numeric_at_min = any(is_numeric_value(v) for v in row_at_min_vals)

            if has_numeric_at_min:
                data_start_r = c_min_r
                data_end_r = c_max_r
                headers = []
                if c_min_r > 1:
                    scan_vals = [str(ws.cell(c_min_r - 1, ci).value or "").strip() for ci in range(c_min_c, c_max_c + 1)]
                    if is_dimension_header_row(scan_vals) or (any(scan_vals) and not any(is_numeric_value(v) for v in scan_vals)):
                        headers = scan_vals

                if not any(headers):
                    m_hdr = sheet_master_headers.get(c_item["sheet"].lower())
                    if m_hdr:
                        m_vals = m_hdr["vals"]
                        m_slice = [m_vals[ci - 1] if 0 <= ci - 1 < len(m_vals) else "" for ci in range(c_min_c, c_max_c + 1)]
                        if any(m_slice):
                            headers = m_slice

                if not any(headers):
                    headers = DEFAULT_DIMENSION_NAMES[:num_cols]
            else:
                headers = [str(v or "").strip() for v in row_at_min_vals]
                data_start_r = c_min_r + 1
                data_end_r = c_max_r

            target_N = data_end_r - data_start_r + 1

            cont_m = [[ws.cell(r, ci).value for ci in range(c_min_c, c_max_c + 1)] for r in range(data_start_r, data_end_r + 1)]

            bps_m = []
            if b_item:
                b_ws = wb[b_item["sheet"]]
                b_min_c, b_min_r = b_item["min_c"], b_item["min_r"]
                b_max_c, b_max_r = b_item["max_c"], b_item["max_r"]
                b_tot_r = b_max_r - b_min_r + 1
                b_start = b_min_r if b_tot_r == target_N else (b_min_r + 1)
                b_end = b_max_r
                bps_m = [[b_ws.cell(r, ci).value for ci in range(b_min_c, b_max_c + 1)] for r in range(b_start, b_end + 1)]

            gw_m = []
            if g_item:
                g_ws = wb[g_item["sheet"]]
                g_min_c, g_min_r = g_item["min_c"], g_item["min_r"]
                g_max_c, g_max_r = g_item["max_c"], g_item["max_r"]
                g_tot_r = g_max_r - g_min_r + 1
                g_start = g_min_r if g_tot_r == target_N else (g_min_r + 1)
                g_end = g_max_r
                gw_m = [[g_ws.cell(r, ci).value for ci in range(g_min_c, g_max_c + 1)] for r in range(g_start, g_end + 1)]

            labels_by_col_idx = {}
            category_display_name = token.upper()

            if l_item:
                l_ws = wb[l_item["sheet"]]
                l_min_c, l_min_r = l_item["min_c"], l_item["min_r"]
                l_max_c, l_max_r = l_item["max_c"], l_item["max_r"]
                l_tot_r = l_max_r - l_min_r + 1
                l_start = l_min_r if l_tot_r == target_N else (l_min_r + 1)
                l_end = l_max_r

                label_matrix = [[str(l_ws.cell(r, ci).value or "").strip() for ci in range(l_min_c, l_max_c + 1)] for r in range(l_start, l_end + 1)]
                if label_matrix and label_matrix[0] and label_matrix[0][0] and label_matrix[0][0].lower() != "none":
                    category_display_name = label_matrix[0][0]

                l_num_cols = l_max_c - l_min_c + 1
                for ci in range(min(num_cols, l_num_cols)):
                    labels_by_col_idx[ci] = [label_matrix[ri][ci] for ri in range(len(label_matrix))]
            else:
                if c_min_c > 1:
                    adj_cols_count = min(num_cols, c_min_c - 1)
                    for idx in range(adj_cols_count):
                        col_vals = [str(ws.cell(r, idx + 1).value or "").strip() for r in range(data_start_r, data_end_r + 1)]
                        labels_by_col_idx[idx] = col_vals
                    if 0 in labels_by_col_idx and labels_by_col_idx[0] and labels_by_col_idx[0][0]:
                        v0 = labels_by_col_idx[0][0]
                        if v0 and v0.lower() != "none":
                            category_display_name = v0.upper()

            columns_data = []
            flat_records = []

            for col_idx, col_name in enumerate(headers):
                if not col_name:
                    col_name = DEFAULT_DIMENSION_NAMES[col_idx] if col_idx < len(DEFAULT_DIMENSION_NAMES) else f"Dim_{col_idx+1}"

                dim_labels = labels_by_col_idx.get(col_idx, [])
                if not dim_labels:
                    dim_labels = DEFAULT_DIMENSION_LABELS.get(col_name, [])

                blocks = []
                for row_idx in range(len(cont_m)):
                    share_raw = cont_m[row_idx][col_idx] if row_idx < len(cont_m) and col_idx < len(cont_m[row_idx]) else None
                    share_num = clean_num(share_raw)
                    if share_num is None:
                        continue

                    share_pct = round(share_num * 100.0, 1) if 0 < share_num <= 1.0 else round(share_num, 1)

                    bps_num = 0
                    if bps_m and row_idx < len(bps_m) and col_idx < len(bps_m[row_idx]):
                        b_raw = bps_m[row_idx][col_idx]
                        b_clean = clean_num(b_raw)
                        bps_num = int(round(b_clean)) if b_clean is not None else 0

                    gw_pct = 0.0
                    if gw_m and row_idx < len(gw_m) and col_idx < len(gw_m[row_idx]):
                        g_raw = gw_m[row_idx][col_idx]
                        g_clean = clean_num(g_raw)
                        if g_clean is not None:
                            gw_pct = round(g_clean * 100.0, 1) if abs(g_clean) <= 5.0 and g_clean != 0 else round(g_clean, 1)

                    if col_idx == 0 or "total" in col_name.lower():
                        actual_label = category_display_name
                        generic_label = category_display_name
                    elif dim_labels and row_idx < len(dim_labels) and dim_labels[row_idx] and dim_labels[row_idx].lower() != "none":
                        actual_label = dim_labels[row_idx]
                        generic_label = f"ITEM {row_idx + 1}"
                    else:
                        actual_label = f"ITEM {row_idx + 1}"
                        generic_label = f"ITEM {row_idx + 1}"

                    is_pos = (bps_num >= 0 if b_item else gw_pct >= 0)
                    blocks.append({
                        "label": actual_label,
                        "actual_label": actual_label,
                        "generic_label": generic_label,
                        "share": share_pct,
                        "growth": gw_pct,
                        "bps": bps_num,
                        "is_positive": is_pos
                    })
                    flat_records.append({
                        "Product": category_display_name,
                        "Column": col_name,
                        "Label": actual_label,
                        "Share %": share_pct,
                        "Growth %": gw_pct,
                        "BPS Impact": bps_num
                    })

                columns_data.append({
                    "column_name": col_name,
                    "blocks": blocks,
                    "total_share": sum(b["share"] for b in blocks)
                })

            results[category_display_name] = {
                "columns_data": columns_data,
                "df": pl.DataFrame(flat_records)
            }

        return results
    except Exception:
        return None


def extract_excel_named_ranges(
    excel_path: Union[str, bytes, io.BytesIO],
    **kwargs: Any
) -> Dict[str, Dict[str, Any]]:
    """
    Parses workbook named ranges with support for Name Manager defined names
    (contribution_*, bps_*, growth_*, label_*).
    Uses ultra-fast selective zip + fastexcel loading with graceful openpyxl fallback.
    Returns dictionary of product_name -> {columns_data, df}.
    """
    # 1. Ultra-fast selective extraction (loads only referenced sheets)
    fast_results = fast_extract_excel_named_ranges(excel_path)
    if fast_results:
        return fast_results

    # 2. Graceful fallback to openpyxl
    openpyxl_results = openpyxl_extract_excel_named_ranges(excel_path)
    if openpyxl_results:
        return openpyxl_results

    return {}



def compute_columnar_decomposition(
    df: Union[pl.DataFrame, pd.DataFrame],
    dimensions: List[str],
    metric_col: str,
    growth_col: Optional[str] = None,
    bps_col: Optional[str] = None,
    benchmark_total: Optional[float] = None
) -> List[Dict[str, Any]]:
    """
    Computes columnar decomposition for any custom generic dataset using high-performance Polars.
    """
    if isinstance(df, pd.DataFrame):
        df = pl.from_pandas(df)

    lower_cols = [c.lower() for c in df.columns]
    if "column" in lower_cols and "label" in lower_cols and "share" in lower_cols:
        col_map = {c.lower(): c for c in df.columns}
        result = []
        ordered_columns = df[col_map["column"]].unique().to_list()
        for col_name in ordered_columns:
            sub = df.filter(pl.col(col_map["column"]) == col_name)
            blocks = []
            for b_idx, row in enumerate(sub.to_dicts()):
                growth_val = row.get(col_map.get("growth", "growth"), 0.0) or 0.0
                bps_val = row.get(col_map.get("bps", "bps"), 0) or 0
                actual_lbl = str(row[col_map["label"]])
                generic_lbl = f"ITEM {b_idx + 1}"
                blocks.append({
                    "label": actual_lbl,
                    "actual_label": actual_lbl,
                    "generic_label": generic_lbl,
                    "share": float(row[col_map["share"]]),
                    "growth": float(growth_val),
                    "bps": int(bps_val),
                    "is_positive": (bps_val >= 0 if col_map.get("bps") else growth_val >= 0)
                })
            result.append({
                "column_name": col_name,
                "blocks": blocks,
                "total_share": sum(b["share"] for b in blocks)
            })
        return result

    result = []
    if not metric_col or metric_col not in df.columns:
        num_cols = [c for c, dt in zip(df.columns, df.dtypes) if dt in (pl.Int8, pl.Int16, pl.Int32, pl.Int64, pl.UInt8, pl.UInt16, pl.UInt32, pl.UInt64, pl.Float32, pl.Float64)]
        if num_cols:
            metric_col = num_cols[0]
        else:
            return []

    total_val = benchmark_total or float(df[metric_col].sum() or 0.0)
    if total_val == 0.0:
        total_val = 1.0

    for dim in dimensions:
        if dim not in df.columns:
            continue

        aggs = [pl.col(metric_col).sum().alias(metric_col)]
        if growth_col and growth_col in df.columns:
            aggs.append(pl.col(growth_col).mean().alias("growth"))
        if bps_col and bps_col in df.columns:
            aggs.append(pl.col(bps_col).sum().alias("bps"))

        grouped = df.group_by(dim).agg(aggs)
        grouped = grouped.with_columns(
            (pl.col(metric_col) / total_val * 100.0).round(1).alias("share")
        )

        if "growth" not in grouped.columns:
            grouped = grouped.with_columns(pl.lit(0.0).alias("growth"))
        else:
            grouped = grouped.with_columns(pl.col("growth").round(1))

        if "bps" not in grouped.columns:
            grouped = grouped.with_columns((pl.col("growth") * 10).cast(pl.Int64).alias("bps"))
        else:
            grouped = grouped.with_columns(pl.col("bps").round(0).cast(pl.Int64))

        grouped = grouped.sort(metric_col, descending=True)

        blocks = []
        for b_idx, row in enumerate(grouped.to_dicts()):
            growth_val = float(row["growth"] or 0.0)
            bps_val = int(row["bps"] or 0)
            is_pos = (bps_val >= 0 if bps_col else growth_val >= 0)
            actual_lbl = str(row[dim])
            generic_lbl = f"ITEM {b_idx + 1}"
            blocks.append({
                "label": actual_lbl,
                "actual_label": actual_lbl,
                "generic_label": generic_lbl,
                "share": float(row["share"]),
                "growth": growth_val,
                "bps": bps_val,
                "is_positive": is_pos,
                "raw_value": float(row[metric_col] or 0.0)
            })

        result.append({
            "column_name": dim,
            "blocks": blocks,
            "total_share": sum(b["share"] for b in blocks)
        })

    return result
